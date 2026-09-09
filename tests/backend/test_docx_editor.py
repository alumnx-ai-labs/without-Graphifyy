import io

from docx import Document

from backend.tools.docx_editor import append_bullets_after, apply_paragraph_replacements


def _make_docx_bytes(paragraphs):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraph_texts(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs]


def test_apply_paragraph_replacements_replaces_matching_text():
    original = _make_docx_bytes(["Built internal tools", "Managed a small team"])

    updated = apply_paragraph_replacements(
        original, {"Built internal tools": "Built and shipped internal developer tools"}
    )

    texts = _paragraph_texts(updated)
    assert "Built and shipped internal developer tools" in texts
    assert "Managed a small team" in texts
    assert "Built internal tools" not in texts


def test_apply_paragraph_replacements_ignores_non_matching_keys():
    original = _make_docx_bytes(["Only paragraph"])

    updated = apply_paragraph_replacements(original, {"Nonexistent text": "New text"})

    assert _paragraph_texts(updated) == ["Only paragraph"]


def test_append_bullets_after_inserts_new_bullets():
    original = _make_docx_bytes(["Experience", "Software Engineer, Acme"])

    updated = append_bullets_after(
        original, "Software Engineer, Acme", ["Shipped feature X", "Reduced latency by 30%"]
    )

    texts = _paragraph_texts(updated)
    idx = texts.index("Software Engineer, Acme")
    assert texts[idx + 1] == "Shipped feature X"
    assert texts[idx + 2] == "Reduced latency by 30%"
