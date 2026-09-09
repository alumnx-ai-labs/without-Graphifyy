import io

from docx import Document
from pypdf import PdfReader

MIN_WELL_FORMED_PARAGRAPHS = 5


def extract_resume_text(filename: str, content: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".txt"):
        return content.decode("utf-8", errors="ignore")
    if lower.endswith(".docx"):
        return _extract_docx_text(content)
    if lower.endswith(".pdf"):
        return _extract_pdf_text(content)
    raise ValueError(f"Unsupported file type: {filename}")


def _extract_docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def is_docx_well_formed(content: bytes) -> bool:
    try:
        doc = Document(io.BytesIO(content))
    except Exception:
        return False
    non_empty_paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    return len(non_empty_paragraphs) >= MIN_WELL_FORMED_PARAGRAPHS
