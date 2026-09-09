import io

from docx import Document


def apply_paragraph_replacements(docx_bytes: bytes, replacements: dict[str, str]) -> bytes:
    doc = Document(io.BytesIO(docx_bytes))
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            new_text = replacements[paragraph.text]
            _set_paragraph_text_preserving_style(paragraph, new_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def append_bullets_after(docx_bytes: bytes, after_paragraph_text: str, bullets: list[str]) -> bytes:
    doc = Document(io.BytesIO(docx_bytes))
    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text == after_paragraph_text:
            anchor = paragraph
            break
    if anchor is None:
        raise ValueError(f"No paragraph found matching: {after_paragraph_text!r}")

    insert_after_element = anchor._p
    for bullet_text in bullets:
        new_paragraph = doc.add_paragraph(bullet_text, style="List Bullet")
        insert_after_element.addnext(new_paragraph._p)
        insert_after_element = new_paragraph._p

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _set_paragraph_text_preserving_style(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    first_run = paragraph.runs[0]
    first_run.text = new_text
    for extra_run in paragraph.runs[1:]:
        extra_run.text = ""
