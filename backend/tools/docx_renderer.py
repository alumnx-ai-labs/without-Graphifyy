import io

from docx import Document

from backend.models import ResumeProfile


def render_profile_to_docx(profile: ResumeProfile) -> bytes:
    doc = Document()

    doc.add_heading(profile.contact.name, level=0)
    contact_line = " | ".join(
        part
        for part in [profile.contact.email, profile.contact.phone, profile.contact.location]
        if part
    )
    if contact_line:
        doc.add_paragraph(contact_line)

    if profile.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(profile.summary)

    if profile.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(profile.skills))

    if profile.experience:
        doc.add_heading("Experience", level=1)
        for entry in profile.experience:
            date_range = " - ".join(
                part for part in [entry.start_date, entry.end_date] if part
            )
            header = f"{entry.title}, {entry.company}"
            if date_range:
                header += f" ({date_range})"
            doc.add_paragraph(header, style="Heading 3")
            for bullet in entry.bullets:
                doc.add_paragraph(bullet, style="List Bullet")

    if profile.education:
        doc.add_heading("Education", level=1)
        for entry in profile.education:
            line = f"{entry.degree}, {entry.institution}"
            if entry.year:
                line += f" ({entry.year})"
            doc.add_paragraph(line)

    if profile.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in profile.certifications:
            doc.add_paragraph(cert, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
